#!/usr/bin/env python3
"""
combine_form.py
Combines signframe.json (skeleton from Signframe) + matriz.xlsx (field mapping)
→ form-definition-final.json (complete form definition).

Usage:
    python3 combine_form.py
    python3 combine_form.py --signframe path/to.json --matriz path/to.xlsx --output path/out.json
"""

import json, sys, os, re, argparse, unicodedata
from collections import OrderedDict

try:
    import openpyxl
except ImportError:
    sys.exit("ERROR: pip install openpyxl")

# ─── Constants ───────────────────────────────────────────────────────────────

NEVER_CONDITION = json.dumps({
    "logic": "and",
    "conditions": [{"fieldId": "field_NEVER_EXISTS", "operator": "not_empty"}]
})

TYPE_MAP = {
    'texto': 'text', 'alfanumerico': 'text', 'alfanumérico': 'text',
    'numerico': 'number', 'numérico': 'number',
    'fecha': 'date', 'combo': 'select', 'lista': 'select',
    'select': 'select', 'dropdown': 'select',
    'radio': 'radio', 'radio/combo': 'radio',
    'checkbox': 'checkbox', 'check': 'checkbox',
    'email': 'email',
    'comentario informativo': 'readonly', 'informativo': 'readonly',
    'titulo': 'heading', 'textarea': 'textarea',
}

NATIVE_TYPE_MAP = {'Tx': 'text', 'Btn': 'checkbox', 'Ch': 'select', 'Sig': 'signature'}

NUMBER_FORMATS = {
    'monto': '#.##0,00', 'entero': '#.##0', 'porcentaje': '00.00',
}

PATH_CORRECTIONS = {
    'datosGenerales.personas': 'datosFormulario.personas',
    'Consentimiento': 'consentimiento',
    'dispotabilidadCarencia': 'disputabilidadCarencia',
}

VALID_OPERATORS = frozenset(['not_empty', 'empty', 'equals'])

# ─── Flexible header detection ───────────────────────────────────────────────

HEADER_PATTERNS = [
    (re.compile(r'^#$'),                                          'rowNum'),
    (re.compile(r'secci[oó]n', re.I),                             'seccionPdf'),
    (re.compile(r'acroform\s*actual', re.I),                      'acroActual'),
    (re.compile(r'acroform\s*propuesto', re.I),                   'acroPropuesto'),
    (re.compile(r'etiqueta', re.I),                               'etiqueta'),
    (re.compile(r'nombre\s*interno|source\s*name', re.I),         'sourceName'),
    (re.compile(r'^tipo$', re.I),                                 'nativeType'),
    (re.compile(r'grupo', re.I),                                  'grupo'),
    (re.compile(r'p[aá]gina', re.I),                              'pagina'),
    (re.compile(r'path.*principal', re.I),                        'pathPrincipal'),
    (re.compile(r'path.*secundari', re.I),                        'pathsSecundarios'),
    (re.compile(r'pre.*rellena', re.I),                           'preRellenado'),
    (re.compile(r'obligatori', re.I),                             'obligatorio'),
    (re.compile(r'max\s*length|largo', re.I),                     'maxLength'),
    (re.compile(r'patr[oó]n|regex', re.I),                        'patron'),
    (re.compile(r'^formato$', re.I),                              'formato'),
    (re.compile(r'visibilidad.*condicional|condicional.*visib', re.I), 'visibilidadCondicional'),
    (re.compile(r'cat[aá]logo.*nombre|nombre.*cat[aá]logo', re.I),    'catalogo'),
    (re.compile(r'opciones.*json|lovable|opciones.*formato', re.I),   'opcionesJson'),
    (re.compile(r'tipo\s*de\s*dato', re.I),                      'tipoDato'),
    (re.compile(r'regla\s*original', re.I),                       'reglaOriginal'),
    (re.compile(r'hoja.*cat[aá]logo', re.I),                     'hojaCatalogo'),
]


# ─── Reading ─────────────────────────────────────────────────────────────────

def _is_yes(val):
    return str(val or '').strip().lower() in ('sí', 'si', 'yes', 'true', '1')


def detect_headers(ws):
    col_map = {}
    for c in range(1, ws.max_column + 1):
        header = str(ws.cell(row=1, column=c).value or '').strip()
        if not header:
            continue
        for pattern, prop in HEADER_PATTERNS:
            if pattern.search(header) and prop not in col_map:
                col_map[prop] = c
                break
    return col_map


def read_matriz(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    col_map = detect_headers(ws)

    print(f"\n  Columnas detectadas en {os.path.basename(path)}:")
    for prop, col in sorted(col_map.items(), key=lambda x: x[1]):
        print(f"    Col {col}: {ws.cell(row=1, column=col).value!r} -> {prop}")

    rows = []
    for r in range(2, ws.max_row + 1):
        row = {}
        for prop, col in col_map.items():
            row[prop] = ws.cell(row=r, column=col).value
        if not row.get('sourceName') and not row.get('acroActual'):
            continue

        row['obligatorio'] = _is_yes(row.get('obligatorio'))
        row['preRellenado'] = _is_yes(row.get('preRellenado'))

        if row.get('maxLength') is not None:
            try:
                row['maxLength'] = int(float(str(row['maxLength'])))
            except (ValueError, TypeError):
                row['maxLength'] = None

        row['optionsParsed'] = None
        if row.get('opcionesJson'):
            try:
                row['optionsParsed'] = json.loads(str(row['opcionesJson']))
            except json.JSONDecodeError:
                pass

        row['_row'] = r
        rows.append(row)

    wb.close()
    print(f"    -> {len(rows)} filas de datos\n")
    return rows


def read_signframe(path):
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def collect_sf_fields(data):
    fields = []
    for sec in data.get('sections', []):
        for sub in sec.get('subsections', []):
            for fld in sub.get('fields', []):
                fields.append(fld)
    return fields


# ─── Helpers ─────────────────────────────────────────────────────────────────

def source_key(field):
    sm = field.get('sourceMeta')
    if sm and sm.get('sourceName'):
        return sm['sourceName']
    fid = field.get('id', '')
    return fid[6:] if fid.startswith('field_') else fid


def to_key(name):
    if not name:
        return 'otros'
    k = unicodedata.normalize('NFD', name.lower())
    k = re.sub(r'[̀-ͯ]', '', k)
    k = re.sub(r'[^a-z0-9]+', '_', k)
    return k.strip('_') or 'otros'


def section_parent(name):
    m = re.match(r'^(.+?)\s+(\d+)\s*$', (name or '').strip())
    if m:
        return m.group(1).strip(), m.group(2)
    return (name or '').strip(), None


def correct_path(path):
    if not path:
        return path
    for wrong, right in PATH_CORRECTIONS.items():
        path = path.replace(wrong, right)
    return path


# ─── Type / format resolution ────────────────────────────────────────────────

def resolve_type(row):
    tipo = str(row.get('tipoDato', '') or '').strip().lower()
    if tipo in TYPE_MAP:
        return TYPE_MAP[tipo]
    native = str(row.get('nativeType', '') or '').strip()
    if native in NATIVE_TYPE_MAP:
        return NATIVE_TYPE_MAP[native]
    if row.get('grupo'):
        return 'radio'
    if row.get('optionsParsed'):
        return 'select'
    return 'text'


def resolve_number_format(row):
    fmt = str(row.get('formato', '') or '').strip().lower()
    if not fmt:
        return None
    if any(k in fmt for k in ('monto', 'moneda', 'currency')):
        return NUMBER_FORMATS['monto']
    if any(k in fmt for k in ('porcentaje', '%')):
        return NUMBER_FORMATS['porcentaje']
    if any(k in fmt for k in ('entero', 'integer')):
        return NUMBER_FORMATS['entero']
    if fmt in ('numérico', 'numerico', 'numero'):
        return NUMBER_FORMATS['entero']
    return None


# ─── Options ──────────────────────────────────────────────────────────────────

def build_options(row):
    opts = row.get('optionsParsed')
    if not opts or not isinstance(opts, list):
        return None
    result = []
    for o in opts:
        if isinstance(o, str):
            result.append({'value': o, 'label': o})
        else:
            result.append({
                'value':    o.get('value', o.get('jsonValue', o.get('label', ''))),
                'label':    o.get('label', o.get('value', '')),
                'jsonValue': o.get('jsonValue', o.get('value', o.get('label', ''))),
                'pdfValue':  o.get('pdfValue', o.get('label', o.get('value', ''))),
            })
    return result


# ─── Conditional visibility ──────────────────────────────────────────────────

def build_conditional_visibility(raw, label_to_id):
    if not raw:
        return None
    trimmed = str(raw).strip()
    if not trimmed:
        return None

    if trimmed[0] in ('{', '['):
        try:
            parsed = json.loads(trimmed)
            _ensure_field_prefix(parsed)
            _validate_operators(parsed)
            return json.dumps(parsed)
        except json.JSONDecodeError:
            pass

    m = re.search(r'[Ss]i\s+"([^"]+)"\s+(?:seleccionado|marcado)', trimmed, re.I)
    if not m:
        m = re.search(r'[Ss]i\s+(?:se\s+)?(?:selecciona|elige|marca)\s+"?([^"]+?)"?\s*(?:->|→|,|\s+mostrar)', trimmed, re.I)
    if m:
        ref_label = m.group(1).strip()
        fid = label_to_id.get(ref_label.lower())
        if not fid:
            ref = re.sub(r'[^a-z0-9_]', '', ref_label.replace(' ', '_').lower())
            fid = 'field_' + ref
        return json.dumps({
            'logic': 'and',
            'conditions': [{'fieldId': fid, 'operator': 'not_empty'}]
        })

    return None


def _ensure_field_prefix(obj):
    if not isinstance(obj, dict):
        return
    for c in obj.get('conditions', []):
        fid = c.get('fieldId', '')
        if fid and not fid.startswith('field_'):
            c['fieldId'] = 'field_' + fid


def _validate_operators(obj):
    if not isinstance(obj, dict):
        return
    for c in obj.get('conditions', []):
        op = c.get('operator', '')
        if op and op not in VALID_OPERATORS:
            c['operator'] = 'not_empty'
        if op == 'equals' and 'value' not in c:
            c['value'] = ''


# ─── Radio group matching ────────────────────────────────────────────────────

def _find_matching_option(label, options):
    ll = label.lower().strip()
    for o in options:
        if o.get('label', '').lower().strip() == ll:
            return o
    for o in options:
        ol = o.get('label', '').lower()
        if ll in ol or ol in ll:
            return o
    return None


def _humanize_group(grupo):
    return str(grupo).replace('_', ' ').title()


# ─── Field enrichment ────────────────────────────────────────────────────────

def enrich_field(field, mx, radio_groups, label_to_id):
    """Enrich a signframe field with matrix data.
    REGLA DE ORO: id and sourceMeta are NEVER modified.
    autoFillConcat, repeaterConfig, sourceMeta are PRESERVED as-is.
    """
    ftype = resolve_type(mx)
    path = correct_path(mx.get('pathPrincipal'))
    options = build_options(mx)
    cond_vis = build_conditional_visibility(mx.get('visibilidadCondicional'), label_to_id)
    num_fmt = resolve_number_format(mx)

    field['label'] = mx.get('etiqueta') or field.get('label', '')
    field['type'] = ftype
    field['required'] = mx.get('obligatorio', False)

    # readOnly: false for any field with sourceMeta that paints the PDF
    if field.get('sourceMeta'):
        field['readOnly'] = False
    elif 'readOnly' not in field:
        field['readOnly'] = False

    if 'hidden' not in field:
        field['hidden'] = False
    if 'width' not in field:
        field['width'] = 'full'

    # prefillMode
    if mx.get('preRellenado') is True:
        field['prefillMode'] = 'required'
    elif mx.get('preRellenado') is False:
        field['prefillMode'] = 'none'

    # Paths
    if path:
        field['salidaJSON'] = path
        field['jsonOutputPath'] = path
        field['prefillKey'] = path
        field['excludeFromJson'] = False
    elif not field.get('salidaJSON'):
        field['excludeFromJson'] = True

    # Secondary paths
    if mx.get('pathsSecundarios'):
        raw = str(mx['pathsSecundarios'])
        sep = '|' if '|' in raw else ','
        paths = [p.strip() for p in raw.split(sep) if p.strip()]
        if paths:
            field['mappedPaths'] = paths

    # conditionalVisibility
    if cond_vis:
        field['conditionalVisibility'] = cond_vis
    elif 'conditionalVisibility' not in field:
        field['conditionalVisibility'] = None

    if 'conditionalRequired' not in field:
        field['conditionalRequired'] = None

    # maxLength
    if mx.get('maxLength'):
        field['maxLength'] = mx['maxLength']

    # validationPattern
    if mx.get('patron'):
        field['validationPattern'] = str(mx['patron'])

    # options
    if options:
        field['options'] = options

    # numberFormat
    if num_fmt:
        field['jsonNumberFormat'] = num_fmt

    # Checkbox: checkedPdfValue depends on sourceMeta presence
    if ftype == 'checkbox':
        if field.get('sourceMeta'):
            field['checkedPdfValue'] = True
            field['checkedJsonValue'] = True
        # No sourceMeta → leave existing or omit (UI helper, can be "X")

    # Width hints
    if ftype == 'date':
        field['width'] = 'half'
    if ftype == 'select' and options:
        field['width'] = 'half'

    # autoFillConcat: PRESERVE, never overwrite
    # repeaterConfig: PRESERVE, never overwrite

    # Radio group handling
    grupo = mx.get('grupo')
    if grupo:
        grupo = str(grupo).strip()
    if grupo and grupo in radio_groups and len(radio_groups[grupo]) > 1:
        field['type'] = 'radio'
        group_rows = radio_groups[grupo]
        group_ids = []
        for gr in group_rows:
            sn = str(gr.get('sourceName', gr.get('acroActual', '')))
            group_ids.append('field_' + sn)

        field['radioGroupLabel'] = _humanize_group(grupo)
        field['radioGroupFields'] = [gid for gid in group_ids if gid != field.get('id')]

        if options:
            field['options'] = options

        etiqueta = mx.get('etiqueta', '')
        if options:
            matched_opt = _find_matching_option(str(etiqueta), options)
            if matched_opt:
                field['jsonValue'] = matched_opt.get('jsonValue', matched_opt.get('value', etiqueta))
                field['pdfValue'] = matched_opt.get('pdfValue', matched_opt.get('label', etiqueta))
            else:
                field['jsonValue'] = etiqueta
                field['pdfValue'] = etiqueta
        else:
            field['jsonValue'] = etiqueta
            field['pdfValue'] = etiqueta

    return field


# ─── Section organization ────────────────────────────────────────────────────

def _assign_order(fields):
    for i, f in enumerate(fields):
        f['order'] = i + 1


def organize_sections(fields_with_section, unmatched_sf):
    # Collect unique section names in order of appearance
    seen = set()
    section_names_ordered = []
    section_fields = {}

    for field, sec_name in fields_with_section:
        if sec_name not in seen:
            seen.add(sec_name)
            section_names_ordered.append(sec_name)
            section_fields[sec_name] = []
        section_fields[sec_name].append(field)

    # Group numbered sections: BENEFICIARIO 1/2/3 → parent BENEFICIARIO
    parent_groups = OrderedDict()
    parent_insertion_order = []
    for sec_name in section_names_ordered:
        base, num = section_parent(sec_name)
        if base not in parent_groups:
            parent_groups[base] = []
            parent_insertion_order.append(base)
        parent_groups[base].append((sec_name, num))

    sections = []
    sec_order = 1

    for parent_name in parent_insertion_order:
        children = parent_groups[parent_name]
        has_numbered = any(num is not None for _, num in children)

        if has_numbered and len(children) > 1:
            subsections = []
            sub_order = 1
            for child_name, _ in children:
                child_fields = section_fields.get(child_name, [])
                _assign_order(child_fields)
                sub_id = 'subsection_' + to_key(child_name)
                subsections.append({
                    'id': sub_id,
                    'title': child_name.strip().title(),
                    'order': sub_order,
                    'fields': child_fields,
                    'childrenOrder': [{'kind': 'field', 'id': f['id']} for f in child_fields],
                })
                sub_order += 1

            sec_id = 'section_' + to_key(parent_name)
            parent_title = parent_name.strip().title()
            if not parent_title.endswith('s') and not parent_title.endswith('es'):
                parent_title += 's'

            sections.append({
                'id': sec_id,
                'title': parent_title,
                'order': sec_order,
                'subsections': subsections,
                'childrenOrder': [{'kind': 'subsection', 'id': s['id']} for s in subsections],
            })
            sec_order += 1
        else:
            for child_name, _ in children:
                child_fields = section_fields.get(child_name, [])
                _assign_order(child_fields)
                sec_id = 'section_' + to_key(child_name)
                sub_id = 'subsection_' + to_key(child_name)
                clean_title = child_name.strip().title()

                subsection = {
                    'id': sub_id,
                    'title': clean_title,
                    'order': 1,
                    'fields': child_fields,
                    'childrenOrder': [{'kind': 'field', 'id': f['id']} for f in child_fields],
                }

                sections.append({
                    'id': sec_id,
                    'title': clean_title,
                    'order': sec_order,
                    'subsections': [subsection],
                    'childrenOrder': [{'kind': 'subsection', 'id': sub_id}],
                })
                sec_order += 1

    # Unmatched signframe fields → hidden system section with NEVER condition
    if unmatched_sf:
        _assign_order(unmatched_sf)
        sub_id = 'subsection_sistema_oculto'
        subsection = {
            'id': sub_id,
            'title': 'Datos del Sistema (oculto)',
            'order': 1,
            'conditionalVisibility': NEVER_CONDITION,
            'fields': unmatched_sf,
            'childrenOrder': [{'kind': 'field', 'id': f['id']} for f in unmatched_sf],
        }
        sections.append({
            'id': 'section_sistema',
            'title': 'Sistema',
            'order': sec_order,
            'subsections': [subsection],
            'childrenOrder': [{'kind': 'subsection', 'id': sub_id}],
        })

    return sections


# ─── Validation ──────────────────────────────────────────────────────────────

def validate_output(sections, warnings):
    total = 0
    for sec in sections:
        for sub in sec.get('subsections', []):
            for f in sub.get('fields', []):
                total += 1
                # order must not be 0
                if f.get('order', 0) == 0:
                    warnings.append(f"FIX: field '{f.get('id')}' tenia order=0, corregido")
                    f['order'] = 1

                # conditionalVisibility validation
                cv = f.get('conditionalVisibility')
                if cv and isinstance(cv, str):
                    try:
                        parsed = json.loads(cv)
                        for cond in parsed.get('conditions', []):
                            op = cond.get('operator', '')
                            if op not in VALID_OPERATORS:
                                old_op = op
                                cond['operator'] = 'not_empty'
                                f['conditionalVisibility'] = json.dumps(parsed)
                                warnings.append(
                                    f"FIX: field '{f.get('id')}' operador invalido '{old_op}' -> 'not_empty'"
                                )
                            if op == 'equals' and 'value' not in cond:
                                cond['value'] = ''
                                f['conditionalVisibility'] = json.dumps(parsed)
                            fid = cond.get('fieldId', '')
                            if fid and not fid.startswith('field_'):
                                cond['fieldId'] = 'field_' + fid
                                f['conditionalVisibility'] = json.dumps(parsed)
                    except json.JSONDecodeError:
                        pass

                # checkbox with sourceMeta must use true, not "X"
                if f.get('type') == 'checkbox' and f.get('sourceMeta'):
                    cpv = f.get('checkedPdfValue')
                    if cpv in ('X', 'x'):
                        f['checkedPdfValue'] = True
                        f['checkedJsonValue'] = True
                        warnings.append(
                            f"FIX: checkbox '{f.get('id')}' con sourceMeta tenia 'X' -> true"
                        )
    return total


# ─── Main ────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description='Combina signframe.json + matriz.xlsx -> form-definition-final.json'
    )
    parser.add_argument('--signframe', default='./signframe.json',
                        help='Path al JSON skeleton de Signframe (default: ./signframe.json)')
    parser.add_argument('--matriz', default='./matriz.xlsx',
                        help='Path al XLSX de mapeo (default: ./matriz.xlsx)')
    parser.add_argument('--output', default='./form-definition-final.json',
                        help='Path de salida (default: ./form-definition-final.json)')
    args = parser.parse_args()

    # ── 1. Read inputs ──
    print(f"Leyendo {args.signframe}...")
    sf_data = read_signframe(args.signframe)

    print(f"Leyendo {args.matriz}...")
    mx_rows = read_matriz(args.matriz)

    sf_fields = collect_sf_fields(sf_data)
    print(f"  {len(sf_fields)} campos en signframe.json")
    print(f"  {len(mx_rows)} filas en matriz\n")

    # ── 2. Build lookups ──
    mx_by_source = {}
    mx_by_acro = {}
    for row in mx_rows:
        sn = row.get('sourceName')
        aa = row.get('acroActual')
        if sn:
            mx_by_source[str(sn)] = row
        if aa:
            mx_by_acro[str(aa)] = row

    # Label → field_id lookup (for conditionalVisibility natural-language parsing)
    label_to_id = {}
    for row in mx_rows:
        etiq = row.get('etiqueta')
        sn = row.get('sourceName', row.get('acroActual', ''))
        if etiq and sn:
            label_to_id[str(etiq).lower().strip()] = 'field_' + str(sn)

    # ── 3. Detect radio groups ──
    # Only groups where members are Btn-type (checkboxes in PDF) become radio groups.
    # Tx-type groups (like fecha_solicitud = dia/mes/ano) are related fields, not radios.
    candidate_groups = {}
    for row in mx_rows:
        g = row.get('grupo')
        if g:
            g = str(g).strip()
            if g not in candidate_groups:
                candidate_groups[g] = []
            candidate_groups[g].append(row)

    radio_groups = {}
    for g, rows in candidate_groups.items():
        has_btn = any(str(r.get('nativeType', '')).strip() == 'Btn' for r in rows)
        has_radio_type = any('radio' in str(r.get('tipoDato', '')).lower() for r in rows)
        if has_btn or has_radio_type:
            radio_groups[g] = rows

    # ── 4. Cross-reference and enrich ──
    fields_with_section = []
    unmatched_sf = []
    warnings = []
    used_mx_keys = set()

    for field in sf_fields:
        key = source_key(field)
        mx = mx_by_source.get(key) or mx_by_acro.get(key)

        if mx:
            used_mx_keys.add(str(mx.get('sourceName', '')))
            used_mx_keys.add(str(mx.get('acroActual', '')))
            sec_name = str(mx.get('seccionPdf') or 'Sin Seccion').strip()
            enrich_field(field, mx, radio_groups, label_to_id)
            fields_with_section.append((field, sec_name))
        else:
            unmatched_sf.append(field)
            warnings.append(
                f"signframe field '{field.get('id')}' (source={key}) "
                f"sin match en matriz -> seccion Sistema"
            )

    # Matrix rows without signframe match
    for row in mx_rows:
        sn = str(row.get('sourceName', ''))
        aa = str(row.get('acroActual', ''))
        if sn not in used_mx_keys and aa not in used_mx_keys:
            warnings.append(
                f"fila de matriz sourceName='{sn}' (acro='{aa}') "
                f"sin match en signframe.json"
            )

    # Repeater warnings
    for field in sf_fields:
        if field.get('repeaterConfig'):
            warnings.append(
                f"Field '{field.get('id')}' tiene repeaterConfig "
                f"-> preservado tal cual, revisar manualmente"
            )

    # ── 5. Organize sections ──
    sections = organize_sections(fields_with_section, unmatched_sf)

    # ── 6. Validate ──
    total_fields = validate_output(sections, warnings)

    # ── 7. Build output ──
    output = dict(sf_data)
    output['sections'] = sections

    with open(args.output, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    # ── 8. Report ──
    n_sections = len(sections)
    n_matched = len(fields_with_section)
    n_unmatched = len(unmatched_sf)

    print(f"\nGenerado: {args.output}")
    print(f"  {n_sections} secciones")
    print(f"  {total_fields} campos totales")
    print(f"  {n_matched} matcheados con matriz")
    print(f"  {n_unmatched} sin match (-> seccion Sistema)")

    if warnings:
        print(f"\n  {len(warnings)} advertencias:")
        for w in warnings:
            print(f"    {w}")

    print("\nDone.")


if __name__ == '__main__':
    main()
